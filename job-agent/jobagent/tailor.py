"""Generate a tailored resume + cover letter for a specific job using an LLM."""
import re
from pathlib import Path

from . import config, db, llm
from .profile import Profile

RESUME_SYSTEM = (
    "You are an expert resume writer. You only reorganize, emphasize, and rephrase "
    "the candidate's REAL experience -- you never invent skills, employers, titles, "
    "or accomplishments that aren't grounded in the source resume."
)

RESUME_PROMPT = """\
Tailor the resume below for the target job. Reorder and rephrase bullet points to
foreground the most relevant experience and keywords for this job. Keep it truthful
and grounded in the original content -- do not fabricate anything. Output clean
Markdown suitable for a one-page resume.

ORIGINAL RESUME:
---
{resume}
---

TARGET JOB:
Title: {title}
Company: {company}
Description:
{description}
---

Output only the tailored resume in Markdown.
"""

COVER_LETTER_SYSTEM = (
    "You are an expert cover letter writer. You write concise, specific, non-generic "
    "cover letters grounded only in the candidate's real background."
)

COVER_LETTER_PROMPT = """\
Write a concise (250-350 word) cover letter for this candidate applying to the job
below. Be specific about why the candidate fits (reference real experience from the
resume), avoid generic filler phrases, and keep the tone confident and direct.

CANDIDATE NAME: {name}
CANDIDATE RESUME:
---
{resume}
---

TARGET JOB:
Title: {title}
Company: {company}
Description:
{description}
---

Output only the cover letter body text (no letterhead, no markdown headers).
"""


def _slugify(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")[:60]


def tailor_job(job_id: str, profile: Profile) -> tuple[str, str]:
    job = db.get_job(job_id)
    if job is None:
        raise ValueError(f"No job found with id '{job_id}'")

    resume_text = profile.resume_text()

    resume_prompt = RESUME_PROMPT.format(
        resume=resume_text[:6000],
        title=job["title"],
        company=job["company"],
        description=(job["description"] or "")[:4000],
    )
    tailored_resume = llm.chat(resume_prompt, system=RESUME_SYSTEM, max_tokens=2000)

    cover_prompt = COVER_LETTER_PROMPT.format(
        name=profile.name or "Candidate",
        resume=resume_text[:6000],
        title=job["title"],
        company=job["company"],
        description=(job["description"] or "")[:4000],
    )
    cover_letter = llm.chat(cover_prompt, system=COVER_LETTER_SYSTEM, max_tokens=800)

    config.ensure_dirs()
    job_dir = config.APPLICATIONS_DIR / f"{_slugify(job['company'])}_{_slugify(job['title'])}_{job_id.replace(':', '-')}"
    job_dir.mkdir(parents=True, exist_ok=True)

    resume_path = job_dir / "resume.md"
    cover_letter_path = job_dir / "cover_letter.md"
    resume_path.write_text(tailored_resume, encoding="utf-8")
    cover_letter_path.write_text(cover_letter, encoding="utf-8")

    _export_docx(tailored_resume, job_dir / "resume.docx")
    _export_docx(cover_letter, job_dir / "cover_letter.docx")

    db.update_tailored_paths(job_id, str(resume_path), str(cover_letter_path))
    return str(resume_path), str(cover_letter_path)


def _export_docx(markdown_text: str, out_path: Path) -> None:
    import docx

    document = docx.Document()
    for line in markdown_text.splitlines():
        stripped = line.strip()
        if not stripped:
            document.add_paragraph("")
        elif stripped.startswith("### "):
            document.add_heading(stripped[4:], level=3)
        elif stripped.startswith("## "):
            document.add_heading(stripped[3:], level=2)
        elif stripped.startswith("# "):
            document.add_heading(stripped[2:], level=1)
        elif stripped.startswith(("- ", "* ")):
            document.add_paragraph(stripped[2:], style="List Bullet")
        else:
            document.add_paragraph(stripped)
    document.save(str(out_path))
